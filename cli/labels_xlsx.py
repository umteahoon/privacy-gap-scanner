# 라벨링용 엑셀 파일 생성 / 엑셀 → CSV 변환
#   생성: python cli/labels_xlsx.py make  [experiment/labels_annotator_sample.csv]
#         → experiment/labels_A.xlsx (annotator1 칸), experiment/labels_B.xlsx (annotator2 칸)
#   변환: python cli/labels_xlsx.py tocsv <입력.xlsx> <출력.csv>   (merge-labels.js 가 자동 호출)
import csv, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

CHOICES = ['명시', '포괄', '미명시']

def make(src='experiment/labels_annotator_sample.csv'):
    rows = list(csv.DictReader(open(src, encoding='utf-8-sig')))
    for who, col in (('A', 'annotator1'), ('B', 'annotator2')):
        wb = Workbook(); ws = wb.active; ws.title = '라벨링'
        head = ['site', 'entity', 'categories', 'hosts', 'policy_url', col, 'note']
        ws.append(head)
        for r in rows:
            ws.append([r['site'], r['entity'], r['categories'], r['hosts'], r['policy_url'], '', ''])
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='335CFF')
        judge = ws.cell(1, 6); judge.fill = PatternFill('solid', fgColor='D12F3A')
        for i in range(2, ws.max_row + 1):
            link = ws.cell(i, 5)
            if link.value: link.hyperlink = link.value; link.font = Font(color='1F4FD8', underline='single')
            ws.cell(i, 6).fill = PatternFill('solid', fgColor='FFF3DF')
            ws.cell(i, 4).alignment = Alignment(wrap_text=False)
        dv = DataValidation(type='list', formula1='"' + ','.join(CHOICES) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle='입력 오류', error='명시 / 포괄 / 미명시 중에서 선택하세요.')
        ws.add_data_validation(dv); dv.add(f'F2:F{ws.max_row}')
        for col_letter, w in zip('ABCDEFG', (22, 24, 26, 40, 46, 12, 30)):
            ws.column_dimensions[col_letter].width = w
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = ws.dimensions
        guide = wb.create_sheet('판정 기준')
        for line in [
            '판정 기준 (해당 사이트의 처리방침 전체를 읽고 판정)', '',
            '명시: 그 사업자 이름이나 서비스명이 적혀 있음 (예: Google Analytics, 카카오픽셀, 크리테오)',
            '포괄: 이름은 없지만 그 기능을 가리키는 표현이 있음 (예: 맞춤형 광고를 위한 행태정보 수집, 웹 로그 분석 도구)',
            '미명시: 이름도, 해당 기능에 대한 고지도 없음', '',
            '주의: "구글"이 간편 로그인 제공자로만 적혀 있고 광고·분석 고지가 없으면, 광고 사업자 Google은 명시가 아님',
            '      (범주 표현이 있으면 포괄, 없으면 미명시. note 칸에 "로그인 맥락만 있음"이라고 적기)', '',
            '다른 평가자의 파일이나 시스템 판정(labels.csv)은 라벨링이 끝날 때까지 보지 마세요.',
        ]:
            guide.append([line])
        guide.column_dimensions['A'].width = 110
        out = f'experiment/labels_{who}.xlsx'
        wb.save(out); print('saved', out, len(rows), 'rows')

def tocsv(src, dst):
    ws = load_workbook(src, data_only=True)['라벨링']
    with open(dst, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            w.writerow(['' if v is None else str(v).strip() for v in row])

if __name__ == '__main__':
    if sys.argv[1] == 'make': make(*sys.argv[2:])
    elif sys.argv[1] == 'tocsv': tocsv(sys.argv[2], sys.argv[3])
